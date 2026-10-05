"""Read a Shopify product export (.csv or .xlsx) into product rows."""

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

METAFIELD_COLUMN = re.compile(
    r"^(?P<label>.*) \(product\.metafields\.(?P<namespace>[\w-]+)\.(?P<key>[\w-]+)\)$"
)
MARKET_COLUMN = re.compile(r"^(?P<kind>Included|Price|Compare At Price) / (?P<market>.+)$")


@dataclass
class Metafield:
    namespace: str
    key: str
    label: str
    column: str


@dataclass
class Export:
    path: Path
    columns: list[str]
    rows: list[dict[str, str]]
    products: list[dict[str, str]] = field(init=False)

    def __post_init__(self) -> None:
        # Shopify writes one row per image/variant; only the first row of a product carries the title.
        self.products = [r for r in self.rows if (r.get("Title") or "").strip()]

    @property
    def metafields(self) -> list[Metafield]:
        found = []
        for col in self.columns:
            if m := METAFIELD_COLUMN.match(col):
                found.append(Metafield(m["namespace"], m["key"], m["label"], col))
        return found

    def metafield_column(self, key: str, namespace: str = "custom") -> str | None:
        for mf in self.metafields:
            if mf.key == key and mf.namespace == namespace:
                return mf.column
        return None

    @property
    def markets(self) -> list[str]:
        seen: list[str] = []
        for col in self.columns:
            if (m := MARKET_COLUMN.match(col)) and m["market"] not in seen:
                seen.append(m["market"])
        return seen


def load_export(path: str | Path) -> Export:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return _load_xlsx(path)
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        rows = [{k: (v or "") for k, v in row.items() if k is not None} for row in reader]
        return Export(path, list(reader.fieldnames or []), rows)


def _load_xlsx(path: Path) -> Export:
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True).active
    it = sheet.iter_rows(values_only=True)
    columns = [str(c) if c is not None else "" for c in next(it)]
    rows = []
    for values in it:
        rows.append({col: _cell(v) for col, v in zip(columns, values, strict=False) if col})
    return Export(path, [c for c in columns if c], rows)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)
